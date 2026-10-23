#!/usr/bin/env python3
"""
Pull fresh Shopify data for the original product ids in plan_variants and merge it into
plan_products / plan_variants. Handles, titles and vendors are never touched.

Setup:   pip install requests openpyxl      (openpyxl only needed for .xlsx)
         export SHOPIFY_STORE=your-store.myshopify.com
         (optional) export TEST_SHOPIFY_STORE=... TEST_SHOPIFY_TOKEN=...   and add --test to use them
         export SHOPIFY_TOKEN=shpat_xxx      (Admin API token: read_products, read_inventory)

Usage:   python pull_shopify_data.py --products plan_products.csv --variants plan_variants.csv --offset 0 --limit 500   (then --offset 1, 2, ...)
         (.csv is fastest; .xlsx also works. Files are updated in place unless --out-suffix is given.)

--limit = how many unique original product ids per run; --offset = which set (0, 1, 2, ...). The start position is
calculated as offset x limit, so limit 500 + offset 2 starts at id 1000. Ids are in variants-file order (skipped rows and
blank ids ignored). --limit 0 = everything (offset ignored).
Speed: batches of 50 ids, 4 parallel calls (--workers). For thousands of products add --bulk: one background job reads the
whole catalog without rate limits, then only the selected ids are used.
"""
import argparse, csv, json, os, sys, time, datetime
from concurrent.futures import ThreadPoolExecutor
import requests

SESSION = requests.Session()   # reuses the connection = faster

API_VERSION = "2025-01"
CAT_PREFIX = "Media > Music & Sound Recordings > "


def full_category(c):
    c = (c or "").strip()
    return c if not c or c.startswith("Media >") else CAT_PREFIX + c


UNITS = {"GRAMS": "g", "KILOGRAMS": "kg", "OUNCES": "oz", "POUNDS": "lb"}
PRODUCT_FIELDS = ["image_url", "product_type", "tags", "genres", "upc", "release_date", "type", "category"]   # all metafields come from the product
VARIANT_ONLY = ["price", "cost_per_item", "quantity", "weight", "weight_unit"]
EXTRA = ["fetch_note", "fetched_at"]

QUERY = """
query($ids:[ID!]!){ nodes(ids:$ids){ ... on Product {
  id productType tags category{ fullName }
  featuredMedia{ ... on MediaImage{ image{ url } } }
  genres: metafield(namespace:"custom", key:"genres"){ value }
  upc: metafield(namespace:"custom", key:"upc"){ value }
  release_date: metafield(namespace:"custom", key:"release_date"){ value }
  ptype: metafield(namespace:"custom", key:"type"){ value }
  variants(first:3){ nodes{ price inventoryQuantity
    inventoryItem{ unitCost{ amount } measurement{ weight{ value unit } } } } }
} } }
"""


# ---------- file io ----------
def read_table(path):
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            r = csv.DictReader(f)
            return list(r.fieldnames), [dict(x) for x in r]
    import openpyxl
    ws = openpyxl.load_workbook(path).active
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    return hdr, [{h: ("" if v is None else v) for h, v in zip(hdr, r)} for r in rows[1:]]


def write_table(path, hdr, rows):
    if path.lower().endswith(".csv"):
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=hdr, extrasaction="ignore")
            w.writeheader()
            w.writerows(rows)
        return
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(hdr)
    for r in rows:
        ws.append([r.get(h, "") if r.get(h, "") != "" else None for h in hdr])
    wb.save(path)


# ---------- shopify ----------
def gql(store, token, query, variables=None):
    url = f"https://{store}/admin/api/{API_VERSION}/graphql.json"
    for attempt in range(6):
        r = SESSION.post(url, json={"query": query, "variables": variables or {}},
                         headers={"X-Shopify-Access-Token": token}, timeout=60)
        if r.status_code in (429, 502, 503):
            time.sleep(2 * (attempt + 1)); continue
        r.raise_for_status()
        j = r.json()
        errs = j.get("errors") or []
        if any("THROTTLED" in json.dumps(e) for e in errs):
            time.sleep(2 * (attempt + 1)); continue
        if errs:
            raise RuntimeError(json.dumps(errs)[:500])
        ts = (j.get("extensions", {}).get("cost", {}) or {}).get("throttleStatus", {})
        if ts and ts.get("currentlyAvailable", 1000) < 150:
            time.sleep(max(0.5, (150 - ts["currentlyAvailable"]) / max(ts.get("restoreRate", 50), 1)))
        return j["data"]
    raise RuntimeError("Shopify kept throttling/failing; try a smaller --batch or fewer --workers")


def fetch_batch(store, token, ids):
    """ids: list of numeric product-id strings -> {id: product dict or None}"""
    gids = [f"gid://shopify/Product/{i}" for i in ids]
    data = gql(store, token, QUERY, {"ids": gids})
    return {gid.rsplit("/", 1)[1]: (node or None) for gid, node in zip(gids, data["nodes"])}


BULK_QUERY = """{ products { edges { node {
  id productType tags category{ fullName }
  featuredMedia{ ... on MediaImage{ image{ url } } }
  genres: metafield(namespace:"custom", key:"genres"){ value }
  upc: metafield(namespace:"custom", key:"upc"){ value }
  release_date: metafield(namespace:"custom", key:"release_date"){ value }
  ptype: metafield(namespace:"custom", key:"type"){ value }
  variants { edges { node { price inventoryQuantity
    inventoryItem{ unitCost{ amount } measurement{ weight{ value unit } } } } } }
} } } }"""


def parse_bulk(lines):
    """JSONL from a bulk operation -> {product id: node shaped like the normal query result}"""
    out = {}
    for line in lines:
        if not line: continue
        o = json.loads(line)
        if "__parentId" in o:
            parent = out.get(o["__parentId"].rsplit("/", 1)[1])
            if parent is not None: parent["variants"]["nodes"].append(o)
        else:
            o["variants"] = {"nodes": []}
            out[o["id"].rsplit("/", 1)[1]] = o
    return out


def fetch_bulk(store, token):
    """Whole catalog in ONE background job (no rate limits). Best when pulling thousands of products."""
    d = gql(store, token, "mutation($q:String!){ bulkOperationRunQuery(query:$q){ bulkOperation{ id } userErrors{ message } } }", {"q": BULK_QUERY})
    ue = d["bulkOperationRunQuery"]["userErrors"]
    if ue: raise RuntimeError("bulk start failed: " + "; ".join(e["message"] for e in ue))
    while True:
        time.sleep(3)
        op = gql(store, token, "{ currentBulkOperation(type: QUERY){ status errorCode objectCount url } }")["currentBulkOperation"]
        print(f"  bulk job: {op['status']} ({op.get('objectCount')} objects)", flush=True)
        if op["status"] == "COMPLETED": break
        if op["status"] in ("FAILED", "CANCELED", "EXPIRED"): raise RuntimeError(f"bulk job {op['status']} {op.get('errorCode')}")
    if not op.get("url"): return {}
    r = SESSION.get(op["url"], stream=True, timeout=300)
    r.raise_for_status()
    return parse_bulk(l for l in r.iter_lines(decode_unicode=True))


def flatten(p):
    """Shopify product node -> flat dict of our columns (first variant used)."""
    mv = lambda k: (p.get(k) or {}).get("value") or ""
    img = (((p.get("featuredMedia") or {}).get("image")) or {}).get("url") or ""
    vs = (p.get("variants") or {}).get("nodes") or []
    v = vs[0] if vs else {}
    inv = v.get("inventoryItem") or {}
    cost = ((inv.get("unitCost") or {}).get("amount")) or ""
    w = ((inv.get("measurement") or {}).get("weight")) or {}
    d = {
        "image_url": img,
        "product_type": p.get("productType") or "",
        "tags": json.dumps(p.get("tags") or []),
        "genres": mv("genres"),
        "upc": mv("upc"),
        "release_date": mv("release_date"),
        "type": mv("ptype"),
        "category": full_category(((p.get("category") or {}).get("fullName")) or ""),
        "price": v.get("price", ""),
        "cost_per_item": cost,
        "quantity": v.get("inventoryQuantity", ""),
        "weight": w.get("value", ""),
        "weight_unit": UNITS.get(w.get("unit"), w.get("unit") or ""),
        "fetch_note": f"{len(vs)} variants in source product; used first" if len(vs) > 1 else "",
    }
    return d


# ---------- main ----------
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--products", required=True)
    ap.add_argument("--variants", required=True)
    ap.add_argument("--offset", type=int, default=0, help="which set to import (0, 1, 2, ...); starts at offset x limit")
    ap.add_argument("--limit", type=int, default=0, help="how many source ids to import (0 = all)")
    ap.add_argument("--batch", type=int, default=50, help="ids per API call")
    ap.add_argument("--workers", type=int, default=4, help="parallel API calls")
    ap.add_argument("--bulk", action="store_true", help="fetch the whole catalog in one background job (fastest for thousands)")
    ap.add_argument("--test", action="store_true", help="use TEST_SHOPIFY_STORE / TEST_SHOPIFY_TOKEN instead of the live store")
    ap.add_argument("--out-suffix", default="", help="write to <name><suffix>.<ext> instead of in place")
    a = ap.parse_args()
    pre = "TEST_" if a.test else ""
    store, token = os.environ.get(pre + "SHOPIFY_STORE"), os.environ.get(pre + "SHOPIFY_TOKEN")
    if not store or not token:
        sys.exit(f"Set {pre}SHOPIFY_STORE and {pre}SHOPIFY_TOKEN environment variables.")
    if a.test: print("TEST STORE:", store)

    vh, vrows = read_table(a.variants)
    ph, prows = read_table(a.products)
    for r in prows:   # make sure every category has the full "Media > Music & Sound Recordings > " prefix
        if "category" in r: r["category"] = full_category(r["category"])
    for c in PRODUCT_FIELDS + VARIANT_ONLY + EXTRA:
        if c not in vh: vh.append(c)
    for c in PRODUCT_FIELDS + EXTRA:
        if c not in ph: ph.append(c)

    # unique source ids in file order
    order, seen = [], set()
    for r in vrows:
        pid = str(r.get("source_product_id", "")).strip()
        if pid and r.get("status") != "skipped" and pid not in seen:
            seen.add(pid); order.append(pid)
    start = a.offset * a.limit if a.limit else 0
    sel = order[start: (start + a.limit) if a.limit else None]
    print(f"{len(order)} source ids in file; importing {len(sel)} (set {a.offset}, starting at #{start})")
    if not sel: return

    data, fetched = {}, datetime.date.today().isoformat()
    try:
        if a.bulk:
            allp = fetch_bulk(store, token)
            data = {pid: allp.get(pid) for pid in sel}
        else:
            chunks = [sel[i:i + a.batch] for i in range(0, len(sel), a.batch)]
            with ThreadPoolExecutor(max_workers=max(1, a.workers)) as ex:
                for n, res in enumerate(ex.map(lambda c: fetch_batch(store, token, c), chunks), 1):
                    data.update(res)
                    print(f"  fetched {min(n * a.batch, len(sel))}/{len(sel)}", flush=True)
    finally:
        # apply whatever was fetched, even if interrupted
        flat = {pid: (flatten(p) if p else None) for pid, p in data.items()}
        for r in vrows:
            pid = str(r.get("source_product_id", "")).strip()
            if pid in flat:
                f = flat[pid]
                if f is None:
                    r["fetch_note"] = "source product not found in Shopify"
                else:
                    r.update(f)
                r["fetched_at"] = fetched
        # product sheet: use the first variant row of each handle as the primary source
        primary = {}
        for r in vrows:
            if r.get("handle") and r.get("status") != "skipped":
                primary.setdefault(r["handle"], r)
        for r in prows:
            pr = primary.get(r.get("handle"))
            if pr and str(pr.get("source_product_id", "")).strip() in flat:
                for c in PRODUCT_FIELDS + EXTRA:
                    r[c] = pr.get(c, "")
        def out(path):
            if not a.out_suffix: return path
            b, e = os.path.splitext(path); return b + a.out_suffix + e
        write_table(out(a.variants), vh, vrows)
        write_table(out(a.products), ph, prows)
        print("saved", out(a.variants), "and", out(a.products))


if __name__ == "__main__":
    main()
